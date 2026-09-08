from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import Workbook

from .db import DB_PATH, connect, initialize
from .import_excel import EXPECTED_HEADERS, INSERT_SQL, normalize, normalize_badge, normalize_drive_file_id

SPREADSHEET_ID = "1SMbMfK-2T5LroHcycjUbf__pwAYQ6wtUHQocl2EoxmU"
SHEET_GIDS = {
    "LIST": "0",
    "OFFICE_MOVEMENT": "1900000001",
    "ADMINISTRATIVE_DOCUMENTS": "1900000002",
    "COMMENDATIONS": "1900000003",
}
ROOT_DIR = Path(__file__).resolve().parent.parent
SNAPSHOT_PATH = ROOT_DIR / "NBPattendance.xlsx"


def _download_csv(gid: str) -> list[list[str]]:
    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=csv&gid={gid}"
    request = Request(url, headers={"User-Agent": "NBPPersonnelLookup/1.0"})
    with urlopen(request, timeout=30) as response:
        text = response.read().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _write_snapshot(sheet_rows: dict[str, list[list[str]]], path: Path = SNAPSHOT_PATH) -> None:
    workbook = Workbook()
    first = True
    for title, rows in sheet_rows.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = title
        for row in rows:
            sheet.append(row)
    workbook.save(path)


def _header_index(headers: list[str]) -> dict[str, int]:
    return {str(value).strip().upper(): index for index, value in enumerate(headers)}


def _cell(values: list[str], index: dict[str, int], header: str) -> str | None:
    position = index.get(header)
    if position is None or position >= len(values):
        return None
    return normalize(values[position])


def _sync_related(connection, table: str, rows: list[list[str]], columns: list[tuple[str, str]]) -> int:
    connection.execute(f"DELETE FROM {table}")
    if not rows:
        return 0
    index = _header_index(rows[0])
    count = 0
    sql_columns = [column for column, _ in columns] + ["source_order"]
    placeholders = ", ".join("?" for _ in sql_columns)
    sql = f"INSERT INTO {table} ({', '.join(sql_columns)}) VALUES ({placeholders})"
    for source_order, values in enumerate(rows[1:], start=1):
        badge = _cell(values, index, "BADGE NUMBER")
        if not badge:
            continue
        record = []
        for column, header in columns:
            if column == "badge_number":
                record.append(normalize_badge(values[index[header]]) if header in index and index[header] < len(values) else None)
            else:
                record.append(_cell(values, index, header))
        record.append(source_order)
        connection.execute(sql, record)
        count += 1
    return count


def _sync_office_movements(connection, rows: list[list[str]]) -> int:
    """Import movement history while preserving both office and camp on each side.

    Expected layout is:
    RECORD ID | BADGE NUMBER | FROM OFFICE | FROM CAMP | TO OFFICE | TO CAMP | POSITION | FROM DATE | TO DATE | REMARKS

    Older/current sheets may still have the sixth header accidentally named FROM OFFICE.
    When that duplicate appears immediately after TO OFFICE, it is treated as TO CAMP.
    """
    connection.execute("DELETE FROM office_movements")
    if not rows:
        return 0

    headers = [str(value).strip().upper() for value in rows[0]]
    positions: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        positions.setdefault(header, []).append(index)

    def at(values: list[str], pos: int | None) -> str | None:
        if pos is None or pos >= len(values):
            return None
        return normalize(values[pos])

    def first(header: str) -> int | None:
        items = positions.get(header, [])
        return items[0] if items else None

    from_office_pos = first("FROM OFFICE")
    from_camp_pos = first("FROM CAMP")
    to_office_pos = first("TO OFFICE")
    to_camp_pos = first("TO CAMP")

    # Compatibility with the present sheet where column F is still labelled FROM OFFICE.
    duplicate_from_office = positions.get("FROM OFFICE", [])
    if to_camp_pos is None and len(duplicate_from_office) > 1:
        candidate = duplicate_from_office[1]
        if to_office_pos is not None and candidate > to_office_pos:
            to_camp_pos = candidate

    sql = """
        INSERT INTO office_movements (
            record_id, badge_number, from_office, from_camp, to_office, to_camp,
            position, from_date, to_date, remarks, source_order
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """
    count = 0
    for source_order, values in enumerate(rows[1:], start=1):
        badge_raw = at(values, first("BADGE NUMBER"))
        badge = normalize_badge(badge_raw) if badge_raw else None
        if not badge:
            continue
        connection.execute(sql, (
            at(values, first("RECORD ID")),
            badge,
            at(values, from_office_pos),
            at(values, from_camp_pos),
            at(values, to_office_pos),
            at(values, to_camp_pos),
            at(values, first("POSITION")),
            at(values, first("FROM DATE")),
            at(values, first("TO DATE")),
            at(values, first("REMARKS")),
            source_order,
        ))
        count += 1
    return count


def sync_google_sheet(db_path: Path | str = DB_PATH) -> dict:
    all_rows = {name: _download_csv(gid) for name, gid in SHEET_GIDS.items()}
    rows = all_rows["LIST"]
    if not rows:
        raise ValueError("Google Sheet returned no LIST data.")

    headers = [str(value).strip().upper() for value in rows[0]]
    column_map = {index: EXPECTED_HEADERS[header] for index, header in enumerate(headers) if header in EXPECTED_HEADERS}
    missing = sorted({"badge_number", "last_name", "first_name"} - set(column_map.values()))
    if missing:
        raise ValueError(f"Google Sheet is missing required columns: {', '.join(missing)}")

    initialize(db_path)
    imported = 0
    with connect(db_path) as connection:
        for source_order, values in enumerate(rows[1:], start=1):
            record = {field: None for field in EXPECTED_HEADERS.values()}
            for index, field in column_map.items():
                if index >= len(values):
                    continue
                value = values[index]
                if field == "badge_number":
                    record[field] = normalize_badge(value)
                elif field == "drive_file_id":
                    record[field] = normalize_drive_file_id(value)
                else:
                    record[field] = normalize(value)
            if not record["badge_number"]:
                continue
            record["source_order"] = source_order
            connection.execute(INSERT_SQL, record)
            imported += 1

        movement_count = _sync_office_movements(connection, all_rows["OFFICE_MOVEMENT"])
        admin_count = _sync_related(
            connection,
            "administrative_documents",
            all_rows["ADMINISTRATIVE_DOCUMENTS"],
            [
                ("record_id", "RECORD ID"),
                ("badge_number", "BADGE NUMBER"),
                ("date_received", "DATE RECEIVED"),
                ("memo_no", "MEMO NO."),
                ("subject_description", "SUBJECT / DESCRIPTION"),
                ("document_from", "FROM"),
                ("remarks", "REMARKS"),
            ],
        )
        commendation_count = _sync_related(
            connection,
            "commendations",
            all_rows["COMMENDATIONS"],
            [
                ("record_id", "RECORD ID"),
                ("badge_number", "BADGE NUMBER"),
                ("date_received", "DATE RECEIVED"),
                ("award_title", "AWARD / TITLE"),
                ("presented_by", "PRESENTED BY"),
                ("remarks", "REMARKS"),
            ],
        )

    _write_snapshot(all_rows)
    return {
        "ok": True,
        "count": imported,
        "movement_count": movement_count,
        "administrative_count": admin_count,
        "commendation_count": commendation_count,
        "synced_at": datetime.now().isoformat(timespec="seconds"),
        "snapshot": str(SNAPSHOT_PATH),
    }
